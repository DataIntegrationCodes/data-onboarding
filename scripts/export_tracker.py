"""
Regenerate data/tracker.json from the Data Onboarding Tracker workbook.

Usage:
    python scripts/export_tracker.py
    python scripts/export_tracker.py --source "D:\\path\\to\\Data Onboarding Tracker.xlsx"

Requires: openpyxl (pip install openpyxl)
"""
import argparse
import datetime
import json
import os

import openpyxl

DEFAULT_SOURCE = (
    r"C:\Users\HRampelwa.INNOWIND\EDF power solutions\Asset Operations - NSight"
    r"\Data Onboarding Tracker.xlsx"
)
DEFAULT_ASSETS_ROOT = (
    r"C:\Users\HRampelwa.INNOWIND\EDF power solutions\Asset Operations - NSight"
)

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(REPO_ROOT, "data", "tracker.json")

CODE_BY_NAME = {
    "Grassridge": "GRAS", "Chaba": "CHAA", "Waainek": "WAAI", "Wesley": "WESL",
    "Dassiesridge": "DASS", "Sankraal": "SANK", "Phezukomoya": "PHEZ", "Coleskop": "COLE",
    "Umsobomvu": "UMSO", "Hartebeeshoek": "HART", "Mooiplaats": "MOOI", "Avondale": "AVON",
    "Mookodi": "MOOK", "Aggeneis": "AGGE", "Nieuwehoop": "NIEU", "Ararat": "ARAR",
}

SINGLE_COL = {"GRAS": 5, "CHAA": 6, "WAAI": 7, "WESL": 8}
TRIPLE_COL = {
    "DASS": 10, "SANK": 13, "PHEZ": 16, "COLE": 19, "UMSO": 22, "HART": 25,
    "MOOI": 28, "AVON": 31, "MOOK": 34, "AGGE": 37, "NIEU": 40, "ARAR": 43,
}

# Project code -> Technical Information folder, relative to the assets root.
DOC_FOLDER_MAP = {
    "DASS": "Vestas Turbine/DASS/Technical Information",
    "SANK": "GoldWind Turbine/SANK/Technical Information",
    "PHEZ": "GoldWind Turbine/PHEZ/Technical Information",
    "COLE": "GoldWind Turbine/COLE/Technical Information",
    "UMSO": "Nordex Turbine/UMSO/Technical Information",
    "HART": "Nordex Turbine/HART/Technical Information",
    "MOOI": "MOOI/Technical Information",
    "AVON": "AVON/Technical Information",
}

DOC_CATEGORY_KEYWORDS = [
    ("network diagram", "Detailed Network Diagram"),
    ("single line diagram", "Single Line Diagram"),
    ("taglist", "Taglist+Review"),
    ("quadrigram", "Quadrigram"),
    ("deployment book", "Quadrigram"),
    ("data model", "Wind/Solar Data Model"),
]


def jval(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    return v


def categorize_folder(folder_name):
    lname = folder_name.lower()
    for keyword, label in DOC_CATEGORY_KEYWORDS:
        if keyword in lname:
            return label
    return folder_name


def scan_doc_completeness(assets_root, code):
    rel = DOC_FOLDER_MAP.get(code)
    if not rel:
        return {}
    base = os.path.join(assets_root, *rel.split("/"))
    if not os.path.isdir(base):
        return {}
    result = {}
    for entry in sorted(os.listdir(base)):
        full = os.path.join(base, entry)
        if not os.path.isdir(full):
            continue
        count = sum(len(files) for _, _, files in os.walk(full))
        label = categorize_folder(entry)
        result[label] = result.get(label, 0) + count
    return result


def project_col_group(code):
    if code in SINGLE_COL:
        return [SINGLE_COL[code]]
    col = TRIPLE_COL[code]
    return [col, col + 1, col + 2]


def first_non_none(ws, row, cols):
    for c in cols:
        v = ws.cell(row=row, column=c).value
        if v is not None:
            return v
    return None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", default=DEFAULT_SOURCE, help="Path to Data Onboarding Tracker.xlsx")
    parser.add_argument("--assets-root", default=DEFAULT_ASSETS_ROOT,
                         help="Path to the Asset Operations - NSight folder (for doc completeness scan)")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.source, data_only=True)

    ws = wb["Onboarding Dashboard"]
    headers = [c.value for c in ws[1]]
    dash_rows = []
    for row in ws.iter_rows(min_row=2, max_row=21):
        vals = [c.value for c in row]
        if vals[1] is None:
            continue
        dash_rows.append(dict(zip(headers, vals)))

    projects = []
    for r in dash_rows:
        name = (r.get("Project") or "").strip()
        code = CODE_BY_NAME.get(name, name[:4].upper())
        projects.append({
            "code": code,
            "portfolio": r.get("Portfolio"),
            "project": name,
            "technology": r.get("Technology"),
            "capacityMW": r.get("Capacity (MW)"),
            "estCOD": jval(r.get("Est. COD")),
            "daysTillCOD": jval(r.get("Days till COD")) if r.get("Days till COD") != "#VALUE!" else None,
            "progression": r.get("Onboarding progression"),
            "interconnectivity": r.get("Interconnectivity \n(RDL & Manu)"),
            "tagReview": r.get("Tag + Breaking \nlist review"),
            "myNsightsOnboarding": r.get("MyNsights Onboarding"),
            "estOnboardingTimeline": r.get("Est. onboarding timeline"),
            "estProductionDate": jval(r.get("Est. production date")),
            "activeOnNsight": r.get("Active on Nsight"),
            "currentPerfMonitoring": r.get("Current Perf monitoring"),
            "comments": r.get("Comments"),
            "occOnboarding": r.get("OCC onboarding"),
            "edcOnboarding": r.get("eDC onboarding"),
            "nsightInternalTraining": r.get("Nsight internal training"),
        })

    tws = wb["Tracker"]
    task_rows = list(range(6, 28))
    tasks = []
    for r in task_rows:
        label = tws.cell(row=r, column=1).value
        team = tws.cell(row=r, column=2).value
        resp_name = tws.cell(row=r, column=3).value
        note = tws.cell(row=r, column=4).value
        if label is None:
            continue
        tasks.append({"row": r, "label": label, "team": team, "responsible": resp_name, "note": note})

    for p in projects:
        code = p["code"]
        cols = project_col_group(code)
        task_values = {}
        for t in tasks:
            task_values[str(t["row"])] = jval(first_non_none(tws, t["row"], cols))
        p["taskCompletion"] = task_values
        p["countryLead"] = jval(first_non_none(tws, 4, cols))
        p["tasksCompletedCount"] = sum(1 for v in task_values.values() if v is True)
        p["tasksTotalCount"] = len(tasks)
        p["docFiles"] = scan_doc_completeness(args.assets_root, code)

    data = {
        "generatedAt": datetime.datetime.now().strftime("%Y-%m-%d"),
        "tasks": tasks,
        "projects": projects,
    }

    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    print(f"Wrote {OUT_PATH}")
    print(f"Projects: {len(projects)}, tasks: {len(tasks)}")


if __name__ == "__main__":
    main()
